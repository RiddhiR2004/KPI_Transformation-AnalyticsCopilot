from __future__ import annotations

import io
import json
import logging
import time
from pathlib import Path
from typing import Any, Dict, List, Optional
import asyncio

from app.services.llm_providers import get_provider
from app.services.llm.langchain_provider import DemoProvider

logger = logging.getLogger("app.services.transcript")
logger.setLevel(logging.INFO)

TRANSCRIPT_ANALYSIS_SYSTEM_PROMPT = """You are a Senior KPI Transformation Consultant. Your role is to analyze a meeting transcript (or combined summaries of a transcript) and extract strategic business insights.

You must return a single JSON object containing exactly the following keys:
{
  "executive_summary": "A brief summary of the transcript discussions and context (100-200 words)",
  "strategic_priorities": ["List of strategic priorities mentioned in the transcript"],
  "business_challenges": ["List of business challenges or issues mentioned in the transcript"],
  "key_decisions": ["List of key decisions made during the meeting"],
  "action_items": ["List of concrete action items agreed upon"],
  "risks_dependencies": ["List of risks or dependencies highlighted"],
  "functional_areas": ["List of functional areas involved or discussed"],
  "mentioned_metrics": ["List of metrics, KPIs, or measurement candidates specifically mentioned in the text"],
  "stakeholders": ["List of key stakeholders or roles mentioned during the meeting"]
}

Avoid placeholders. Ensure all extracted items are professional, clear, and business-focused.
"""

CHUNK_SUMMARY_SYSTEM_PROMPT = """You are a Senior Business Analyst. Summarize the following section of a meeting transcript.
Extract and highlight:
- Major discussion points
- Key decisions
- Action items
- Stated priorities and challenges
- Mentioned metrics or KPI candidates
- Stakeholders involved

Keep the summary concise and focused on strategic outcomes.
"""


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        return content.decode("utf-8", errors="ignore")
    
    elif ext == ".docx":
        import docx
        doc = docx.Document(io.BytesIO(content))
        return "\n".join([p.text for p in doc.paragraphs])
    
    elif ext == ".pdf":
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(content))
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    
    else:
        raise ValueError("Unsupported file format. Only .txt, .docx, and .pdf are supported.")


def chunk_text(text: str, chunk_size: int = 8000, overlap: int = 1000) -> List[str]:
    """Chunks text into smaller pieces with overlap."""
    if len(text) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
        
    return chunks


async def analyze_transcript_text(raw_text: str) -> Dict[str, Any]:
    provider = get_provider()
    
    # Handle DemoProvider fallback
    if isinstance(provider, DemoProvider):
        logger.info("Using DemoProvider mock insights for transcript analysis")
        return {
            "executive_summary": "Demo Mode: Mock summary of meeting discussing operational issues and system upgrades.",
            "strategic_priorities": [
                "Reduce machine downtime by 20%",
                "Improve inventory tracking accuracy",
                "Enhance cross-departmental collaboration"
            ],
            "business_challenges": [
                "Manual logging causing data lag",
                "High scrap rate in the assembly line",
                "Siloed communication between IT and Operations"
            ],
            "key_decisions": [
                "Approved budget for inventory tracking pilot",
                "Decided to implement real-time machine monitoring system"
            ],
            "action_items": [
                "IT team to review API endpoints for ERP integration by Friday",
                "Schedule a vendor demo for OEE tracking software"
            ],
            "risks_dependencies": [
                "Legacy hardware compatibility issues",
                "Operator training lag and resistance to new tools"
            ],
            "functional_areas": ["Production", "Supply Chain", "IT Services"],
            "mentioned_metrics": ["Downtime Hours", "Scrap Rate", "OEE", "Inventory Accuracy"],
            "stakeholders": ["COO", "Plant Manager", "IT Director", "Finance Controller"]
        }

    # Size Protection check
    MAX_CHAR_THRESHOLD = 12000
    if len(raw_text) > MAX_CHAR_THRESHOLD:
        logger.info(f"Transcript length ({len(raw_text)} chars) exceeds threshold. Starting chunked summarization...")
        chunks = chunk_text(raw_text, chunk_size=8000, overlap=1000)
        
        # Summarize each chunk asynchronously
        tasks = []
        for i, chunk in enumerate(chunks):
            user_prompt = f"=== TRANSCRIPT CHUNK {i+1}/{len(chunks)} ===\n{chunk}"
            tasks.append(
                provider.generate_json(
                    CHUNK_SUMMARY_SYSTEM_PROMPT,
                    user_prompt,
                    step_name=f"summarize_transcript_chunk_{i+1}"
                )
            )
        
        chunk_results = await asyncio.gather(*tasks, return_exceptions=True)
        
        summaries = []
        for idx, res in enumerate(chunk_results):
            if isinstance(res, Exception):
                logger.error(f"Error summarizing chunk {idx+1}: {res}")
                continue
            
            # Extract summary text from response payload
            if isinstance(res, dict):
                summary_text = res.get("summary") or res.get("summary_text") or json.dumps(res)
                summaries.append(summary_text)
            else:
                summaries.append(str(res))
                
        # Merge summaries
        combined_summaries_text = "\n\n".join(
            f"--- Chunk {i+1} Summary ---\n{s}" for i, s in enumerate(summaries)
        )
        
        logger.info("Chunk summaries compiled. Extracting final structured insights...")
        payload = await provider.generate_json(
            TRANSCRIPT_ANALYSIS_SYSTEM_PROMPT,
            combined_summaries_text,
            step_name="analyze_transcript_combined"
        )
        return payload
    else:
        logger.info(f"Transcript length ({len(raw_text)} chars) is within limits. Extracting insights directly...")
        payload = await provider.generate_json(
            TRANSCRIPT_ANALYSIS_SYSTEM_PROMPT,
            raw_text,
            step_name="analyze_transcript"
        )
        return payload
